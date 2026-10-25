from pathlib import Path
import csv
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def flatten_record(record: dict, rubric: dict) -> dict:
    grade = record.get("grade", record)
    row = {
        "ID": grade.get("student_id", ""),
        "Name": grade.get("student_name", ""),
    }
    for task in rubric["tasks"]:
        assessed = next((t for t in grade.get("tasks", []) if t.get("task_id") == task["id"]), {})
        row[f'{task["id"]} Marks'] = assessed.get("marks_awarded", "")
        row[f'{task["id"]} Reason'] = assessed.get("reasoning", "")
    row["Total"] = grade.get("total_marks", "")
    row["Manual Review"] = "YES" if grade.get("manual_review") else "NO"
    row["Review Reasons"] = "; ".join(grade.get("review_reasons", []))
    return row


def export_csv(records: list[dict], rubric: dict, path: str | Path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = [flatten_record(r, rubric) for r in records]
    headers = list(rows[0].keys()) if rows else ["ID", "Name", "Total", "Manual Review", "Review Reasons"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def export_xlsx(records: list[dict], rubric: dict, path: str | Path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    rows = [flatten_record(r, rubric) for r in records]
    headers = list(rows[0].keys()) if rows else ["ID", "Name", "Total", "Manual Review", "Review Reasons"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, header in enumerate(headers, start=1):
        values = [str(ws.cell(row, col_idx).value or "") for row in range(1, ws.max_row + 1)]
        width = min(max(max((len(v) for v in values), default=10) + 2, 12), 48)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    audit = wb.create_sheet("Audit")
    audit_headers = ["ID", "Name", "Task ID", "Marks", "Maximum", "Reasoning", "Confidence", "Evidence", "Manual Review", "Review Reason", "Verifier Approved"]
    audit.append(audit_headers)
    for record in records:
        grade = record.get("grade", record)
        verifier_tasks = {t.get("task_id"): t for t in record.get("verifier", {}).get("tasks", [])}
        for task in grade.get("tasks", []):
            vt = verifier_tasks.get(task.get("task_id"), {})
            audit.append([
                grade.get("student_id", ""), grade.get("student_name", ""),
                task.get("task_id", ""), task.get("marks_awarded", ""),
                task.get("maximum_marks", ""), task.get("reasoning", ""),
                task.get("confidence", ""), json.dumps(task.get("evidence", []), ensure_ascii=False),
                bool(task.get("manual_review")), task.get("review_reason", ""),
                vt.get("approved", ""),
            ])
    for cell in audit[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    for col_idx in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(col_idx)].width = min(
            max(max((len(str(audit.cell(r, col_idx).value or "")) for r in range(1, audit.max_row + 1)), default=10) + 2, 12),
            48
        )
    wb.save(path)
